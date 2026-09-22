"""Load lexicons and attach frequency / length features."""

from __future__ import annotations  # postponed annotation evaluation

import math  # log10 for the Zipf conversion
import urllib.request  # download SUBTLEX files
from pathlib import Path  # filesystem paths

import pandas as pd  # table handling
from wordfreq import zipf_frequency  # external Zipf frequencies (English), used where a corpus value is absent

DATA_DIR = Path(__file__).resolve().parents[1] / "data"  # repo-root/data
MORPHOLEX_XLSX = DATA_DIR / "external" / "MorphoLEX_en.xlsx"  # source MorphoLex workbook
MORPHOLEX_WORDS = DATA_DIR / "morpholex_words.csv"           # built MorphoLex working list
LADEC_CSV = DATA_DIR / "external" / "LADECv1-2019.csv"       # source LADEC file
LADEC_COMPOUNDS = DATA_DIR / "ladec_compounds.csv"          # built LADEC working list
SUBTLEX_PATH = DATA_DIR / "external" / "SUBTLEXus74286.txt"  # cached SUBTLEX-US download
SUBTLEX_GLOVE = DATA_DIR / "subtlex_glove.csv"              # GloVe-filtered SUBTLEX-US subsample
SUBTLEX_NOUNS_GLOVE = DATA_DIR / "subtlex_nouns_glove.csv"  # POS-filtered variant (when --pos is used)
SUBTLEX_URL = (  # SUBTLEX-US word-list download (tab-separated)
    "https://raw.githubusercontent.com/cltl/python-for-text-analysis/master/"
    "Data/SUBTLEX-US/SUBTLEXus74286wordstextversion.txt"
)
SUBTLEX_GR_PATH = DATA_DIR / "external" / "SUBTLEX-GR_restricted.txt"  # cached Greek SUBTLEX download
SUBTLEX_GR_GLOVE = DATA_DIR / "subtlex_gr_glove.csv"                   # GloVe-filtered Greek subsample
SUBTLEX_GR_URL = "https://www.bcbl.eu/sites/default/files/files/SUBTLEX-GR_restricted.txt"  # Greek source
_GREEK_WORD = r"^[\u0370-\u03ff\u1f00-\u1fff]+$"  # regex: all-Greek-letter token (basic + extended ranges)
SUBTLEX_POS_PATH = DATA_DIR / "external" / "SUBTLEX-US_POS_Zipf.xlsx"  # POS+Zipf spreadsheet cache
SUBTLEX_POS_URL = "https://osf.io/download/55d4847a8c5e4a5fe4a6c8d2/"  # POS spreadsheet download


def load_word_csv(path: Path | None = None) -> pd.DataFrame:
    """Load a word CSV and attach length / Zipf / POS features."""
    path = path or MORPHOLEX_WORDS  # default to the MorphoLex working list
    df = pd.read_csv(path)
    df["word"] = df["word"].str.strip().str.lower()  # normalize surface form
    df = df.drop_duplicates("word").reset_index(drop=True)  # one row per word
    df = df[df["word"].str.fullmatch(r"[a-z]+")].copy()  # keep pure a–z words only
    if "n_morphemes" in df.columns:  # coerce morpheme count to numeric if present
        df["n_morphemes"] = pd.to_numeric(df["n_morphemes"], errors="coerce")
    df = df.drop(columns=[c for c in ("class", "is_monomorph") if c in df.columns], errors="ignore")  # legacy cols
    df["length"] = df["word"].str.len()  # character length
    if "zipf_freq" not in df.columns:  # if the CSV lacks a corpus Zipf, fall back to wordfreq
        df["zipf_freq"] = [zipf_frequency(w, "en") for w in df["word"]]
    df["log_freq"] = df["zipf_freq"]  # Zipf is already log10(freq) + constant; kept as an alias
    df = attach_subtlex_pos(df)  # add a dominant-POS column (used for optional coloring/reference)
    return df.reset_index(drop=True)


def build_morpholex(
    xlsx: Path | None = None,   # source workbook (defaults to MORPHOLEX_XLSX)
    out: Path | None = None,    # output CSV (defaults to MORPHOLEX_WORDS)
) -> pd.DataFrame:
    """All alphabetic types from MorphoLex-en (Sánchez-Gutiérrez et al. 2018).

    Sheets are Prefix-Root-Suffix signatures. Inflected forms are dropped when
    the stem is also in the list. All-caps acronyms and rare title-case names
    are dropped. No POS filter and no morphological-class subsample.
    """
    from openpyxl import load_workbook  # local import: only needed when (re)building

    xlsx = xlsx or MORPHOLEX_XLSX
    out = out or MORPHOLEX_WORDS
    if not xlsx.exists():  # the workbook is not shipped in the repo
        raise FileNotFoundError(
            f"MorphoLex workbook not found at {xlsx}. "
            "Download MorphoLEX_en.xlsx from https://github.com/hugomailhot/MorphoLex-en"
        )

    wb = load_workbook(xlsx, read_only=True, data_only=True)  # open read-only, values (not formulas)
    sheets = [s for s in wb.sheetnames if s[0].isdigit()]  # PRS-signature sheets are named like "1-1-0"
    rows = []  # accumulate per-word dicts
    for name in sheets:
        n_pref, n_root, n_suff = map(int, name.split("-"))  # parse the P-R-S counts from the sheet name
        ws = wb[name]
        it = ws.iter_rows(values_only=True)  # row iterator yielding tuples of cell values
        header = list(next(it))              # first row = column headers
        idx = {h: i for i, h in enumerate(header) if h}  # header name → column index
        if "Word" not in idx:  # skip non-data sheets (e.g. "All roots")
            continue
        for r in it:  # remaining rows are words
            if r is None or r[idx["Word"]] is None:  # skip empty rows
                continue
            orig = str(r[idx["Word"]]).strip()  # raw surface form
            word = orig.lower()                  # lowercased form
            if orig.isupper() or not word.isalpha() or len(word) < 3:  # drop acronyms, non-alpha, very short
                continue
            if orig.istitle() and zipf_frequency(word, "en") < 4.0:  # drop rare proper-noun-like Title forms
                continue
            nm = r[idx["Nmorph"]] if "Nmorph" in idx else None  # morpheme count column if available
            rows.append(
                {
                    "word": word,
                    "n_morphemes": int(nm) if nm is not None else n_pref + n_root + n_suff,  # count or P+R+S
                    "n_pref": n_pref,   # #prefixes (from sheet name)
                    "n_root": n_root,   # #roots
                    "n_suff": n_suff,   # #suffixes
                    "prs": str(r[idx["PRS_signature"]]) if "PRS_signature" in idx else name,  # PRS string
                    "note": str(r[idx["MorphoLexSegm"]]) if "MorphoLexSegm" in idx else name,  # segmentation
                }
            )
    wb.close()

    df = pd.DataFrame(rows).drop_duplicates("word")  # NOTE: keeps the FIRST sheet's parse per duplicate word
    stems = set(df["word"])  # set of all surface forms, for the inflection check below
    inflected = set()        # words judged to be inflected variants of a listed stem
    for w in df["word"]:
        dropped = False
        for suf in ("s", "es", "ed", "ing"):  # common inflectional endings
            if w.endswith(suf) and w[: -len(suf)] in stems:  # e.g. "walks"→"walk" present ⇒ inflected
                inflected.add(w)
                dropped = True
                break
        if not dropped and w.endswith("ing") and (w[:-3] + "e") in stems:  # "baking"→"bake" (drop-e case)
            inflected.add(w)
    df = df.loc[~df["word"].isin(inflected)].copy()  # remove inflected forms whose stem is present
    df["zipf_freq"] = [zipf_frequency(w, "en") for w in df["word"]]  # wordfreq Zipf (MorphoLex has no corpus freq)
    df = df.sort_values("word").reset_index(drop=True)  # deterministic ordering

    keep = ["word", "n_morphemes", "note", "n_pref", "n_root", "n_suff", "prs", "zipf_freq"]  # columns to persist
    out.parent.mkdir(parents=True, exist_ok=True)
    df[keep].to_csv(out, index=False)  # write the working list
    return df


def load_morpholex(path: Path | None = None, rebuild: bool = False) -> pd.DataFrame:
    """Return the MorphoLex working list, building it from the workbook if needed."""
    path = path or MORPHOLEX_WORDS
    if rebuild or not path.exists():  # (re)build only when asked or when the CSV is missing
        build_morpholex(out=path)
    return load_word_csv(path)  # load + attach features


def build_ladec_compounds(src: Path | None = None, out: Path | None = None) -> pd.DataFrame:
    """Closed compounds from LADEC (Gagné, Spalding & Schmidtke 2019).

    correctParse=yes, letters only, native SUBTLEX Zipfvalue required.
    """
    src = src or LADEC_CSV
    out = out or LADEC_COMPOUNDS
    if not src.exists():  # the source CSV is not shipped
        raise FileNotFoundError(
            f"LADEC not found at {src}. Download LADECv1-2019.csv "
            "(Gagné, Spalding & Schmidtke 2019)."
        )
    raw = pd.read_csv(src)
    df = raw.copy()
    df["word"] = df["stim"].astype(str).str.strip().str.lower()  # the compound surface form
    yes = df["correctParse"].astype(str).str.lower().isin(["yes", "1", "true"])  # keep validated parses
    df = df.loc[yes & df["word"].str.fullmatch(r"[a-z]+")].copy()  # validated + pure a–z
    df = df.dropna(subset=["ratingcmp"]).drop_duplicates("word")  # need the human rating; one row per word
    if "Zipfvalue" not in df.columns:  # native SUBTLEX Zipf column is required (no wordfreq fallback)
        raise ValueError("LADEC file has no Zipfvalue column")
    df["zipf_freq"] = pd.to_numeric(df["Zipfvalue"], errors="coerce")  # native corpus Zipf
    df = df.loc[df["zipf_freq"].notna()].copy()  # drop items lacking a native frequency
    df["n_morphemes"] = 2  # every LADEC item is a two-constituent compound by construction
    c1 = df["c1"] if "c1" in df.columns else ""  # first constituent (for the note)
    c2 = df["c2"] if "c2" in df.columns else ""  # second constituent
    df["note"] = [  # human-readable provenance: "c1+c2 ratingcmp=NN.N"
        f"{a}+{b} ratingcmp={r:.1f}"
        for a, b, r in zip(c1, c2, df["ratingcmp"])
    ]
    keep = ["word", "n_morphemes", "note", "ratingcmp", "zipf_freq"]  # base columns
    extra = [c for c in ("c1", "c2") if c in df.columns]  # keep constituents if present
    out.parent.mkdir(parents=True, exist_ok=True)
    df[keep + extra].to_csv(out, index=False)  # write the working list
    return df[keep + extra]


def load_ladec_compounds(path: Path | None = None, rebuild: bool = False) -> pd.DataFrame:
    """Return the LADEC working list, rebuilding from source if needed."""
    path = path or LADEC_COMPOUNDS
    if rebuild or not path.exists():  # rebuild on demand or when missing
        build_ladec_compounds(out=path)
    return load_word_csv(path)  # load + attach features


def download_subtlex(path: Path | None = None, url: str = SUBTLEX_URL) -> Path:
    """Download the SUBTLEX-US word list once and cache it locally."""
    path = path or SUBTLEX_PATH
    if path.exists():  # already cached
        return path
    path.parent.mkdir(parents=True, exist_ok=True)
    req = urllib.request.Request(url, headers={"User-Agent": "morph-opacity-pipeline"})  # UA header
    with urllib.request.urlopen(req, timeout=120) as src, path.open("wb") as dst:  # stream to disk
        dst.write(src.read())
    return path


def download_subtlex_pos(path: Path | None = None, url: str = SUBTLEX_POS_URL) -> Path:
    """Download the SUBTLEX-US POS+Zipf spreadsheet once and cache it."""
    path = path or SUBTLEX_POS_PATH
    if path.exists():
        return path
    path.parent.mkdir(parents=True, exist_ok=True)
    req = urllib.request.Request(url, headers={"User-Agent": "morph-opacity-pipeline"})
    with urllib.request.urlopen(req, timeout=180) as src, path.open("wb") as dst:
        dst.write(src.read())
    return path


def load_subtlex_pos_table(path: Path | None = None) -> pd.DataFrame:
    """Dominant CLAWS POS per SUBTLEX-US type (Brysbaert, New & Keuleers 2012)."""
    cache = DATA_DIR / "cache" / "subtlex_pos.csv"  # small parsed cache of (word, pos)
    if cache.exists():
        return pd.read_csv(cache)  # reuse the parsed cache
    xlsx = download_subtlex_pos(path)  # ensure the spreadsheet is present
    raw = pd.read_excel(xlsx, sheet_name=0, usecols=["Word", "Dom_PoS_SUBTLEX"])  # only two columns needed
    raw = raw.rename(columns={"Word": "word", "Dom_PoS_SUBTLEX": "pos"})
    raw["word"] = raw["word"].astype(str).str.strip().str.lower()  # normalize words
    raw = raw.loc[raw["word"].str.fullmatch(r"[a-z]+")].drop_duplicates("word")  # a–z, unique
    raw["pos"] = raw["pos"].where(raw["pos"].notna(), "unknown").astype(str)  # fill NaN POS
    raw.loc[raw["pos"].isin(["nan", "#N/A", "None", "NaN"]), "pos"] = "unknown"  # normalize junk labels
    cache.parent.mkdir(parents=True, exist_ok=True)
    raw[["word", "pos"]].to_csv(cache, index=False)  # persist the parsed cache
    return raw[["word", "pos"]]


def attach_subtlex_pos(df: pd.DataFrame) -> pd.DataFrame:
    """Left-join a dominant-POS column onto `df` by word."""
    pos = load_subtlex_pos_table()  # (word, pos) table
    out = df.merge(pos, on="word", how="left", suffixes=("", "_subtlex"))  # attach pos
    if "pos_subtlex" in out.columns:  # if df already had a pos column, reconcile the two
        out["pos"] = out["pos"].combine_first(out["pos_subtlex"])  # prefer existing, fill from table
        out = out.drop(columns=["pos_subtlex"])
    out["pos"] = out["pos"].fillna("unknown")  # words not in the POS table
    return out


def load_subtlex_us(
    path: Path | None = None,   # cached word-list path
    min_zipf: float = 0.0,      # drop types below this Zipf
    pos: str | None = None,     # optional dominant-POS filter (e.g. "Noun")
) -> pd.DataFrame:
    """SUBTLEX-US types (Brysbaert & New 2009), letters only.

    If `pos` is set (e.g. \"Noun\"), load the POS+Zipf spreadsheet (Brysbaert,
    New & Keuleers 2012) and keep rows whose dominant CLAWS tag matches.
    Zipf is the official SUBTLEX Zipf-value when POS is used, otherwise
    log10(SUBTLWF)+3.
    """
    if pos:  # POS-filtered branch uses the spreadsheet (which carries its own Zipf-value)
        xlsx = download_subtlex_pos()
        raw = pd.read_excel(xlsx, sheet_name=0)
        raw = raw.rename(columns={"Word": "word", "Zipf-value": "zipf_freq"})  # standardize columns
        raw["word"] = raw["word"].astype(str).str.strip().str.lower()
        dom = raw["Dom_PoS_SUBTLEX"].astype(str)  # dominant POS column
        df = raw.loc[dom.str.fullmatch(pos, case=False)].copy()  # keep matching POS (case-insensitive)
        df = df.loc[df["word"].str.fullmatch(r"[a-z]+")].drop_duplicates("word")  # a–z, unique
        df["zipf_freq"] = pd.to_numeric(df["zipf_freq"], errors="coerce")  # numeric Zipf
        df["note"] = f"subtlex-us {pos.lower()}"  # provenance tag
    else:  # default branch: the plain word-frequency list
        path = download_subtlex(path)
        raw = pd.read_csv(path, sep="\t")  # tab-separated
        raw = raw.rename(columns={"Word": "word"})
        raw["word"] = raw["word"].astype(str).str.strip().str.lower()
        df = raw.loc[raw["word"].str.fullmatch(r"[a-z]+")].copy()  # a–z only
        df = df.drop_duplicates("word")
        df["zipf_freq"] = df["SUBTLWF"].map(lambda x: math.log10(max(float(x), 1e-12)) + 3.0)  # WF/million→Zipf
        df["note"] = "subtlex-us"
    df = df.loc[df["zipf_freq"].notna() & (df["zipf_freq"] >= min_zipf)].copy()  # apply the frequency floor
    df = df.sort_values(["zipf_freq", "word"], ascending=[False, True])  # most frequent first (stable)
    df["length"] = df["word"].str.len()  # character length
    df["n_morphemes"] = pd.NA            # SUBTLEX has no morphological parse
    df["log_freq"] = df["zipf_freq"]     # alias
    df = attach_subtlex_pos(df)          # add dominant POS
    keep = [
        "word",
        "zipf_freq",
        "length",
        "pos",
        "n_morphemes",
        "note",
        "log_freq",
    ]
    return df[keep].reset_index(drop=True)


def download_subtlex_gr(path: Path | None = None, url: str = SUBTLEX_GR_URL) -> Path:
    """Download the Greek SUBTLEX file once and cache it."""
    path = path or SUBTLEX_GR_PATH
    if path.exists():
        return path
    path.parent.mkdir(parents=True, exist_ok=True)
    req = urllib.request.Request(url, headers={"User-Agent": "morph-opacity-pipeline"})
    with urllib.request.urlopen(req, timeout=180) as src, path.open("wb") as dst:
        dst.write(src.read())
    return path


def load_subtlex_gr(path: Path | None = None, min_zipf: float = 0.0) -> pd.DataFrame:
    """SUBTLEX-GR types (Dimitropoulou et al. 2010), Greek letters only.

    Zipf is log10(SUBTLEX_WF)+3 from frequency per million.
    """
    path = download_subtlex_gr(path)  # ensure the file is present
    raw = pd.read_csv(path, sep="\t", skiprows=4, quotechar='"')  # header preamble is 4 lines
    raw.columns = [c.strip().strip('"') for c in raw.columns]  # clean quoted/space-padded headers
    df = raw.rename(columns={"Word": "word"}).copy()
    df["word"] = df["word"].astype(str).str.strip().str.casefold()  # casefold (proper Greek lowercasing)
    df = df.loc[df["word"].str.fullmatch(_GREEK_WORD)].drop_duplicates("word")  # Greek-letter words, unique
    wf = pd.to_numeric(df["SUBTLEX_WF"], errors="coerce")  # frequency per million
    df["zipf_freq"] = wf.map(lambda x: math.log10(max(float(x), 1e-12)) + 3.0)  # convert to Zipf
    df = df.loc[df["zipf_freq"].notna() & (df["zipf_freq"] >= min_zipf)].copy()  # frequency floor
    df["length"] = df["word"].str.len()  # length in characters
    df["note"] = "subtlex-gr"            # provenance tag
    df["log_freq"] = df["zipf_freq"]     # alias
    keep = ["word", "zipf_freq", "length", "note", "log_freq"]
    return df[keep].reset_index(drop=True)
